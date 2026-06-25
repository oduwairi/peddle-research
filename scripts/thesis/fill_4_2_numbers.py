"""Chapter IV §4.2 placeholder fill — swap X.XXX placeholders for measured numbers.

Targets paragraphs in §4.2.2 (Agent-Integrated Variants), §4.2.3 (Per-Platform
Scores), §4.2.4 (Per-Head Scores), and §4.2.5 (Predictor Reliability). All
placeholder values come from:
  - docs/research/SCORING_PREDICTOR_PHASE2_RESULTS_2026-05.md (per-config
    composite and per-head means, per-platform means)
  - docs/research/RQ2_OFFLINE_2x2_RESULTS_2026-05.md (B_pipe, C_pipe means)
  - Figure 4.4 caption (ECE, top/bottom-tier AUCs)
  - §4.1.1 (per-platform n: facebook 93, pinterest 28, reddit 39, tiktok 28,
    twitter 27)

Placeholders span multiple runs, so we rebuild each target paragraph by
concatenating all runs, applying the rewrite, then loading the result into
run[0] and clearing the rest. Run[0] formatting is preserved.

Idempotent — paragraphs whose new text is already present are skipped.

Run with:  uv run python scripts/thesis/fill_4_2_numbers.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

DOCX = Path("docs/research/THESIS.docx")


# Each entry: (anchor_prefix, new_paragraph_text)
# anchor_prefix is the first ~40 chars of the target paragraph (matches the
# concatenated-run text, so it tolerates run-split boundaries).
REWRITES: list[tuple[str, str]] = [
    # §4.2.2 Agent-Integrated Variants (paragraph 477)
    (
        "B_pipe and C_pipe are the agent-integrated",
        (
            "B_pipe and C_pipe are the agent-integrated versions of the flat models. "
            "The models are wrapped in a full agent loop with tool calls such as web "
            "search, URL scraping, similar-page search, and some outputs like image "
            "generation or campaign schema output. The goal is to test whether "
            "integrating the model into this agent flow provides real value over the "
            "base models. First, the agent-integrated systems run through the same "
            "test set that the raw models used. After inferencing the models we "
            "obtained composite means: B_pipe = 0.586, C_pipe = 0.607. Compared to "
            "the single-shot models, we see B (0.611) → B_pipe (0.586), drop of "
            "0.025 composite (4.1% relative). C (0.651) → C_pipe (0.607), drop of "
            "0.044 composite. The decrease is consistent for both cases but hurts "
            "the fine-tuned model more than the base model. We can also see that "
            "C_pipe still beats out B_pipe, similar to the single-shot case, "
            "meaning the agent-wrapped, fine-tuned model still outperforms the "
            "base model. The best single-shot model is ahead of the corresponding "
            "agentic model (C_pipe = 0.607) by 0.044 composite."
        ),
    ),
    # §4.2.3 Per-Platform Scores (paragraph 481)
    (
        "Now focusing on our fine-tuned model's performance across platforms",
        (
            "Now focusing on our fine-tuned model's performance across platforms, "
            "we can see that our fine-tuned model yields composite scores per "
            "platform: pinterest 0.711, twitter 0.659, tiktok 0.652, facebook "
            "0.638, reddit 0.634, with n: facebook 93, pinterest 28, reddit 39, "
            "tiktok 28, twitter 27. Our model has the best performance on the "
            "Pinterest platform, and the worst relative performance on Reddit. "
            "Most notably, our fine-tuned model beats out both the frontier model "
            "configuration as well as the base model on every platform. This shows "
            "that the model internalized platform-native structure where other "
            "models have failed to do so, and that the model is consistently "
            "better than others. Wins are as large as (C +0.074 over B, +0.077 "
            "over A) on Pinterest. Smallest absolute wins are on TikTok (C +0.025 "
            "over B, +0.038 over A). We can also see that the gap to the gold ads "
            "is: tiktok 0.014 (closest to ceiling), pinterest 0.024, reddit "
            "0.027, facebook 0.033, twitter 0.067 (furthest from ceiling)."
        ),
    ),
    # §4.2.4 Per-Head Scores (paragraph 485)
    (
        "We have looked at the composite scores",
        (
            "We have looked at the composite scores. Looking at the three "
            "component heads (§3.7) that sit beneath the composite can also give "
            "insight into the model's individual performance. For the "
            "survivability head, which represents the score for the longevity of "
            "the ad (the Kaplan–Meier survival curve target from the v3 scorer, "
            "§3.3), we can see results of GOLD=0.674, C=0.646, B=0.588, A=0.571, "
            "where C beats B by +0.058 and A by +0.075 on survivability, which is "
            "the biggest margin across the three heads. Gap to GOLD on "
            "survivability = 0.028, which is the smallest gap to the high "
            "baseline across the heads. For the engagement volume head, which "
            "represents the total engagement (reactions, comments, shares) that "
            "a post is likely to get, we see scores of GOLD=0.722, C=0.681, "
            "B=0.651, A=0.650. The core insight is that our fine-tuned model "
            "beats out both configurations by a considerable margin (~+0.030), "
            "while A and B sit essentially tied on this head. Finally, the "
            "engagement velocity head represents the engagement-per-time proxy. "
            "We see scores of GOLD=0.650, C=0.620, B=0.620, A=0.631. It is "
            "interesting to note that C is actually slightly behind A on velocity "
            "(−0.011). However, the margin is quite small. This is the single "
            "case where the fine-tuned model loses out to the frontier model by a "
            "slim margin. Looking at the overall picture, we see that the overall "
            "composite score lift that our fine-tuned model sees is because of "
            "the survivability score (+0.058 over B) and then engagement volume "
            "(+0.030 over B). For reference, gold ads score the highest on all "
            "three heads, which is a signal that our pipeline is representative "
            "of the actual performance."
        ),
    ),
    # §4.2.5 Predictor Reliability (paragraph 489)
    (
        "Given these results, it is important to address the reliability",
        (
            "Given these results, it is important to address the reliability of "
            "the trained predictor (§3.7). Since the predictor is a trained "
            "regressor model on real performance data, its predictions are "
            "grounded. However, it is still trained on a moderate-sized 55k "
            "v3-scored AdFlex corpus, which means drift may happen. The composite "
            "reliability performance on the held-out test split scored Spearman ρ "
            "= 0.722, which is a strong correlation but not perfect. Calibration "
            "error (ECE) = 0.0074 — close to zero, which indicates that the "
            "model predicts scores that match quantile rates. The model's "
            "top-tier AUC = 0.865 and bottom-tier AUC = 0.874 indicate that the "
            "model can reliably separate high-tier performing ads from low-tier "
            "ads. It is worth noting that the Reddit slice of the corpus is the "
            "weakest, since the absence of engagement data from the API makes "
            "the engagement heads have limited data to train on for this "
            "platform. It is important to note that the predictor sees text "
            "only, not the image or video creative, which are part of the "
            "engagement, but for the purposes of this research, we are "
            "fine-tuning a copywriting model, so the trained score accurately "
            "represents the skill. Finally, the trained model is based on our "
            "proxy v3 scorer (§3.3), which means any bias or inaccuracies in "
            "the original score are transferred to the trained model."
        ),
    ),
]


def replace_paragraph_text(paragraph, new_text: str) -> None:
    """Set paragraph.text to new_text by loading it into run[0] and clearing
    runs[1:]. Preserves run[0]'s formatting; downstream runs that carried
    different formatting are flattened (acceptable here — these are plain
    prose paragraphs)."""
    runs = list(paragraph.runs)
    if not runs:
        return
    runs[0].text = new_text
    for r in runs[1:]:
        r.text = ""


def main() -> int:
    doc = Document(str(DOCX))
    applied = 0
    skipped = 0
    for anchor, new_text in REWRITES:
        target = None
        for p in doc.paragraphs:
            joined = "".join(r.text for r in p.runs)
            if joined.startswith(anchor):
                target = p
                current = joined
                break
        if target is None:
            print(f"  MISSING: paragraph starting {anchor!r}")
            continue
        if current.strip() == new_text.strip():
            print(f"  skip (already current): {anchor[:50]!r}")
            skipped += 1
            continue
        replace_paragraph_text(target, new_text)
        print(f"  applied: {anchor[:50]!r}")
        applied += 1
    doc.save(str(DOCX))
    print(f"\nDone. Applied {applied}, skipped {skipped}.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
